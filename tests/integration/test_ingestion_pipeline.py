"""Phase 2 end-to-end: API -> RabbitMQ -> worker -> Docling -> chunk -> embed -> Qdrant.

Uses real Qdrant, Postgres, RabbitMQ and S3 (SeaweedFS) and real Docling (DOCX/TXT, which
need no model download). The dense/sparse embedders are hashing test doubles: the real
models are exercised by the acceptance run (scripts/phase2_check.py).
"""

import asyncio
import io
import time
import uuid

import docx

from beai_rag_core.enums import JobKind
from beai_rag_core.tenancy import TenantScope
from tests.fakes import HashingEmbedder
from tests.integration.conftest import (
    Api,
    Infra,
    Tenant,
    drain_queue,
    new_tenant,
    running_api,
    running_worker,
)

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def returns_policy_docx() -> bytes:
    d = docx.Document()
    d.add_heading("Acme returns policy", level=1)
    d.add_heading("Section 3 – Shipping", level=2)
    d.add_paragraph("Orders ship within two business days. Tracking is emailed.")
    d.add_heading("Section 4 – Returns", level=2)
    d.add_paragraph("Returns are accepted within 30 days, unworn, with receipt.")
    d.add_paragraph("Refunds go back to the original payment method within 5 days.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


async def wait_status(api: Api, t: Tenant, source_id: str, want: str, limit_s: float = 60) -> dict:  # type: ignore[type-arg]
    deadline = time.monotonic() + limit_s
    while time.monotonic() < deadline:
        sources = (await api.client.get(f"{t.base}/sources", headers=t.headers)).json()
        match = [s for s in sources if s["id"] == source_id]
        if want == "gone" and not match:
            return {}
        if match and match[0]["status"] == want:
            return match[0]  # type: ignore[no-any-return]
        if match and match[0]["status"] == "error" and want != "error":
            raise AssertionError(f"source failed: {match[0]['error']}")
        await asyncio.sleep(0.2)
    raise TimeoutError(f"source {source_id} did not reach {want}")


async def upload(api: Api, t: Tenant, name: str, data: bytes, mime: str) -> str:
    r = await api.client.post(
        f"{t.base}/sources/files", headers=t.headers, files=[("files", (name, data, mime))]
    )
    assert r.status_code == 202, r.text
    return str(r.json()[0]["source_id"])


async def test_upload_to_ready_with_citation_payload_then_delete(migrated: Infra) -> None:
    t = new_tenant()
    async with running_api(migrated) as api, running_worker(api):
        started = time.monotonic()
        sid = await upload(api, t, "returns-policy.docx", returns_policy_docx(), DOCX)
        src = await wait_status(api, t, sid, "ready")
        assert time.monotonic() - started < 120
        assert src["chunks"] >= 2 and src["language"] == "en" and src["error"] is None

        scope = TenantScope(t.tenant_id, t.agent_id)
        store = api.resources.store
        records = await store.scroll(scope, source_id=uuid.UUID(sid))
        assert len(records) == src["chunks"]
        payloads = [r.payload or {} for r in records]
        for p in payloads:
            assert p["tenant_id"] == str(t.tenant_id) and p["agent_id"] == str(t.agent_id)
            assert p["file_name"] == "returns-policy.docx" and p["source_type"] == "file"
            assert (
                p["source_id"] == sid and p["lang"] == "en" and p["embed_model"] == "test-hashing@1"
            )
            assert {
                "page",
                "page_end",
                "section",
                "chunk_index",
                "text",
                "knowledge_base_id",
            } <= set(p)
        returns = [p for p in payloads if "30 days, unworn, with receipt" in p["text"]]
        assert len(returns) == 1
        assert returns[0]["section"] == "Acme returns policy > Section 4 – Returns"

        kb = (await api.client.get(f"{t.base}/knowledge-base", headers=t.headers)).json()
        assert kb["status"] == "ready" and kb["kb_version"] >= 2

        # Delete -> all points, the stored file and the row go away.
        assert (
            await api.client.delete(f"{t.base}/sources/{sid}", headers=t.headers)
        ).status_code == 202
        await wait_status(api, t, sid, "gone")
        assert await store.count(scope, source_id=uuid.UUID(sid)) == 0
        assert await store.count(scope) == 0
        kb = (await api.client.get(f"{t.base}/knowledge-base", headers=t.headers)).json()
        assert kb["status"] == "empty"


async def test_chunk_size_change_reindexes_blue_green(migrated: Infra) -> None:
    t = new_tenant()
    text = " ".join(
        f"Sentence number {i} talks about topic {i % 7} in some detail." for i in range(400)
    )
    async with running_api(migrated) as api, running_worker(api) as worker:
        sid = await upload(api, t, "long.txt", text.encode(), "text/plain")
        first = await wait_status(api, t, sid, "ready")
        scope = TenantScope(t.tenant_id, t.agent_id)
        old_ids = {r.id for r in await api.resources.store.scroll(scope, limit=1000)}
        calls_before = worker.dense.document_calls

        r = await api.client.put(
            f"{t.base}/knowledge-base", headers=t.headers, json={"chunk_size": 200}
        )
        assert r.json()["reindexing"] is True
        await asyncio.sleep(0.3)
        second = await wait_status(api, t, sid, "ready")
        assert second["chunks"] > first["chunks"]  # smaller chunks -> more of them
        new = await api.resources.store.scroll(scope, limit=1000)
        assert len(new) == second["chunks"]  # old points removed after the new ones were written
        assert not old_ids & {r.id for r in new}
        assert worker.dense.document_calls > calls_before


async def test_tenant_isolation_in_vector_store(migrated: Infra) -> None:
    a, b = new_tenant(), new_tenant()
    b = Tenant(b.tenant_id, a.agent_id)  # same agent id, different tenant
    async with running_api(migrated) as api, running_worker(api):
        sa = await upload(api, a, "a.txt", b"Tenant A returns: 30 days.", "text/plain")
        sb = await upload(api, b, "b.txt", b"Tenant B returns: 60 days.", "text/plain")
        await wait_status(api, a, sa, "ready")
        await wait_status(api, b, sb, "ready")
        store = api.resources.store
        ra = await store.scroll(TenantScope(a.tenant_id, a.agent_id))
        rb = await store.scroll(TenantScope(b.tenant_id, b.agent_id))
        assert {(r.payload or {})["source_id"] for r in ra} == {sa}
        assert {(r.payload or {})["source_id"] for r in rb} == {sb}


class FlakyEmbedder(HashingEmbedder):
    def __init__(self, failures: int) -> None:
        super().__init__(1024)
        self.failures = failures

    async def embed_documents(self, texts):  # type: ignore[no-untyped-def]
        if self.failures > 0:
            self.failures -= 1
            raise ConnectionError("embedding server unavailable")
        return await super().embed_documents(texts)


async def test_transient_failure_retried_then_succeeds(migrated: Infra) -> None:
    t = new_tenant()
    async with running_api(migrated) as api, running_worker(api, dense=FlakyEmbedder(failures=2)):
        sid = await upload(api, t, "a.txt", b"Retry me please, this is a test.", "text/plain")
        src = await wait_status(api, t, sid, "ready", limit_s=30)
        assert src["error"] is None


async def test_persistent_failure_dead_lettered_with_message(migrated: Infra) -> None:
    t = new_tenant()
    async with running_api(migrated) as api, running_worker(api, dense=FlakyEmbedder(failures=99)):
        sid = await upload(api, t, "a.txt", b"This will never embed.", "text/plain")
        src = await wait_status(api, t, sid, "error", limit_s=60)
        assert "Temporary problem while indexing" in src["error"]
    import aio_pika

    connection = await aio_pika.connect(migrated.rabbitmq_url)
    async with connection:
        channel = await connection.channel()
        dead = await channel.get_queue("ingest.dead")
        found = False
        while (msg := await dead.get(no_ack=True, fail=False)) is not None:
            found = found or sid.encode() in msg.body
        assert found


async def test_unreadable_document_fails_permanently_without_retries(migrated: Infra) -> None:
    t = new_tenant()
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:  # valid OOXML container, invalid document
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("word/document.xml", "<not-really-word/>")
    async with running_api(migrated) as api, running_worker(api) as worker:
        sid = await upload(api, t, "broken.docx", buf.getvalue(), DOCX)
        src = await wait_status(api, t, sid, "error", limit_s=30)
        assert src["error"]
        assert worker.dense.document_calls == 0
    await drain_queue(migrated.rabbitmq_url, JobKind.INDEX)


async def test_business_info_indexed_as_one_page_document(migrated: Infra) -> None:
    t = new_tenant()
    form = {
        "company_name": "Acme Shoes",
        "hours": "Mon-Fri 9-17",
        "contact": "help@acme.test",
        "faq": [{"question": "Do you ship abroad?", "answer": "Yes, to 30 countries."}],
    }
    async with running_api(migrated) as api, running_worker(api):
        r = await api.client.put(f"{t.base}/sources/business-info", headers=t.headers, json=form)
        sid = r.json()["source_id"]
        src = await wait_status(api, t, sid, "ready")
        assert src["pages"] == 1
        records = await api.resources.store.scroll(TenantScope(t.tenant_id, t.agent_id))
        assert {(r.payload or {})["file_name"] for r in records} == {"Business info"}
        assert {(r.payload or {})["page"] for r in records} == {1}
