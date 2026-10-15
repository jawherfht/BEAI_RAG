"""Parsers that need no model download: plain text, business info, Docling DOCX/HTML/MD."""

from pathlib import Path

import docx
import pytest

from worker.parsers.base import ElementKind, ParsedDocument
from worker.parsers.business_info import render_business_info
from worker.parsers.docling_parser import DoclingParser, parse_plain_text

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def test_plain_text_paragraphs() -> None:
    d = parse_plain_text("First para.\nstill first.\n\nSecond para.")
    assert [e.text for e in d.elements] == ["First para.\nstill first.", "Second para."]
    assert all(e.page == 1 for e in d.elements)


def test_business_info_render() -> None:
    d = render_business_info(
        {
            "company_name": "Acme",
            "hours": "9-5",
            "address": "",
            "contact": "hi@acme.test",
            "faq": [{"question": "Do you ship?", "answer": "Yes, worldwide."}],
        }
    )
    assert d.title == "Business info"
    texts = [e.text for e in d.elements]
    assert "Company name: Acme" in texts and "Contact: hi@acme.test" in texts
    assert not any(t.startswith("Address") for t in texts)
    assert "Q: Do you ship?\nA: Yes, worldwide." in texts


def test_parsed_document_json_roundtrip() -> None:
    d = render_business_info({"company_name": "Acme"})
    assert ParsedDocument.from_json(d.to_json()).to_json() == d.to_json()


@pytest.fixture(scope="module")
def parser() -> DoclingParser:
    return DoclingParser(ocr=False)


def test_docling_docx_headings_paragraphs_tables(tmp_path: Path, parser: DoclingParser) -> None:
    document = docx.Document()
    document.add_heading("Returns policy", level=1)
    document.add_heading("Section 4 - Returns", level=2)
    document.add_paragraph("Items can be returned within 30 days, unworn, with receipt.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Item", "Window"
    table.cell(1, 0).text, table.cell(1, 1).text = "Shoes", "30 days"
    path = tmp_path / "policy.docx"
    document.save(str(path))

    d = parser.parse(path, DOCX_MIME)
    kinds = [(e.kind, e.text) for e in d.elements]
    assert (ElementKind.HEADING, "Returns policy") in kinds
    assert (ElementKind.HEADING, "Section 4 - Returns") in kinds
    assert any(k is ElementKind.PARAGRAPH and "30 days, unworn" in t for k, t in kinds)
    tables = [t for k, t in kinds if k is ElementKind.TABLE]
    assert len(tables) == 1 and "| Shoes" in tables[0]
    levels = {e.text: e.level for e in d.elements if e.kind is ElementKind.HEADING}
    assert levels["Returns policy"] < levels["Section 4 - Returns"]


def test_docling_html_and_markdown(tmp_path: Path, parser: DoclingParser) -> None:
    (tmp_path / "p.html").write_text(
        "<html><body><h1>FAQ</h1><p>Shipping takes 3 days.</p></body></html>"
    )
    (tmp_path / "p.md").write_text("# Guide\n\n## Setup\n\nRun the installer.\n")
    h = parser.parse(tmp_path / "p.html", "text/html")
    m = parser.parse(tmp_path / "p.md", "text/markdown")
    assert any(e.text == "Shipping takes 3 days." for e in h.elements)
    assert [e.text for e in m.elements if e.kind is ElementKind.HEADING] == ["Guide", "Setup"]


def test_infrastructure_errors_are_not_blamed_on_the_document() -> None:
    import httpx

    from worker.parsers.docling_parser import is_infrastructure_error

    try:
        try:
            raise httpx.ProxyError("403 Forbidden")
        except httpx.ProxyError as inner:
            raise RuntimeError("model load failed") from inner
    except RuntimeError as outer:
        assert is_infrastructure_error(outer)
    assert not is_infrastructure_error(ValueError("bad pdf"))
