"""Build a knowledge base's documents and questions from its facts module.

    python eval/datasets/build_kb.py kb2_saas_fr [--questions-only]

The facts module (eval/datasets/<kb>/facts.py) defines TITLES {file: title}, FACTS (answerable,
one sentence each, with a paraphrased question and an evidence phrase), UNANSWERABLE (questions
the documents do not answer) and TRICKY (answerable; codes, names, numbers, multi-part).
Every gold label is exact because documents are generated from the facts.
"""

from __future__ import annotations

import html
import importlib.util
import json
import sys
import textwrap
from collections import OrderedDict
from datetime import UTC, datetime
from pathlib import Path
from types import ModuleType
from typing import Any

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))
from make_returns_policy_pdf import build as build_pdf  # type: ignore[import-not-found]  # noqa: E402


def load(kb: str) -> ModuleType:
    spec = importlib.util.spec_from_file_location(f"{kb}_facts", HERE / kb / "facts.py")
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def grouped(facts: list[Any], doc: str) -> OrderedDict[tuple[int | None, str], list[Any]]:
    out: OrderedDict[tuple[int | None, str], list[Any]] = OrderedDict()
    for f in facts:
        if f.doc == doc:
            out.setdefault((f.page, f.section), []).append(f)
    return out


def write_doc(folder: Path, doc: str, title: str, facts: list[Any], rtl: bool) -> None:
    groups = grouped(facts, doc)
    suffix = Path(doc).suffix
    path = folder / doc
    if suffix == ".pdf":
        pages: dict[int, list[tuple[str, int, str]]] = {}
        for (page, section), items in groups.items():
            assert page is not None
            lines = pages.setdefault(page, [("H", 18, title)] if page == 1 else [])
            lines.append(("H", 13, section))
            for line in textwrap.wrap(" ".join(f.sentence for f in items), 95):
                lines.append(("P", 10, line))
        path.write_bytes(build_pdf([pages[p] for p in sorted(pages)]))
    elif suffix == ".docx":
        import docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        d = docx.Document()
        heading = d.add_heading(title, level=1)
        paragraphs = [heading]
        for (_, section), items in groups.items():
            paragraphs.append(d.add_heading(section, level=2))
            paragraphs.append(d.add_paragraph(" ".join(f.sentence for f in items)))
        if rtl:
            for p in paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        d.core_properties.author = "BEAI eval"
        d.core_properties.created = d.core_properties.modified = datetime(2026, 1, 1, tzinfo=UTC)
        d.save(str(path))
    elif suffix == ".md":
        parts = [f"# {title}\n"]
        parts += [f"## {s}\n\n" + " ".join(f.sentence for f in items) + "\n" for (_, s), items in groups.items()]
        path.write_text("\n".join(parts), encoding="utf-8")
    elif suffix == ".html":
        body = "".join(
            f"<h2>{html.escape(s)}</h2>\n<p>{html.escape(' '.join(f.sentence for f in items))}</p>\n"
            for (_, s), items in groups.items()
        )
        lang = ' lang="ar" dir="rtl"' if rtl else ""
        path.write_text(
            f"<!DOCTYPE html>\n<html{lang}><head><meta charset=\"utf-8\"><title>{html.escape(title)}</title></head>"
            f"<body><main><h1>{html.escape(title)}</h1>\n{body}</main></body></html>\n",
            encoding="utf-8",
        )
    elif suffix == ".txt":
        parts = [title] + [" ".join(f.sentence for f in items) for items in groups.values()]
        path.write_text("\n\n".join(parts) + "\n", encoding="utf-8")
    else:
        raise ValueError(f"unsupported document type: {doc}")


def validate(m: ModuleType) -> None:
    """Gold labels are exact only if every evidence phrase is really in its document."""
    for f in m.FACTS:
        assert f.doc in m.TITLES, f
        assert f.evidence in f.sentence, f
        if f.doc.endswith(".pdf"):
            assert f.page is not None, f
            f.sentence.encode("cp1252")  # the raw PDF writer uses WinAnsi Helvetica
    for t in getattr(m, "TRICKY", []):
        text = " ".join(f.sentence for f in m.FACTS if (f.doc, f.page) in set(t.gold))
        for e in t.evidence:
            assert e in text, (t.question, e)
    questions = [f.question for f in m.FACTS] + list(m.UNANSWERABLE)
    assert len(questions) == len(set(questions)), "duplicate question"


def build(kb: str, questions_only: bool = False) -> None:
    m = load(kb)
    validate(m)
    folder = HERE / kb
    rtl = getattr(m, "RTL", False)
    if not questions_only:
        for doc, title in m.TITLES.items():
            write_doc(folder, doc, title, m.FACTS, rtl)
    prefix = getattr(m, "PREFIX", kb.split("_")[0])
    rows: list[dict[str, Any]] = []
    for i, f in enumerate(m.FACTS, start=1):
        rows.append({"id": f"{prefix}-a{i:03d}", "type": "answerable", "question": f.question,
                     "gold": [{"file_name": f.doc, "page": f.page}], "evidence": [f.evidence],
                     "answer": f.sentence})
    for i, q in enumerate(m.UNANSWERABLE, start=1):
        rows.append({"id": f"{prefix}-u{i:03d}", "type": "unanswerable", "question": q, "gold": [], "evidence": []})
    for i, t in enumerate(getattr(m, "TRICKY", []), start=1):
        rows.append({"id": f"{prefix}-t{i:03d}", "type": "tricky", "question": t.question,
                     "gold": [{"file_name": d, "page": p} for d, p in t.gold], "evidence": list(t.evidence)})
    with (folder / "questions.jsonl").open("w", encoding="utf-8") as out:
        for row in rows:
            out.write(json.dumps(row, ensure_ascii=False) + "\n")
    counts = {k: sum(r["type"] == k for r in rows) for k in ("answerable", "unanswerable", "tricky")}
    print(f"{kb}: {len(m.TITLES)} documents, questions {counts}")


if __name__ == "__main__":
    build(sys.argv[1], questions_only="--questions-only" in sys.argv)
