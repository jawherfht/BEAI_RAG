"""Build KB #1 documents and questions from facts.py.

    python eval/datasets/kb1_retail/build.py

Writes returns-policy.pdf, shipping-guide.pdf, product-care.docx, store-faq.md, warranty.txt
and questions.jsonl (type "answerable", gold file/page/section and evidence phrase).
"""

from __future__ import annotations

import json
import sys
import textwrap
from collections import OrderedDict
from pathlib import Path

HERE = Path(__file__).parent
sys.path[:0] = [str(HERE), str(HERE.parent)]

from facts import FACTS, Fact  # type: ignore[import-not-found]  # noqa: E402
from make_returns_policy_pdf import build as build_pdf  # type: ignore[import-not-found]  # noqa: E402

TITLES = {
    "returns-policy.pdf": "Acme Retail – Returns and Shipping Policy",
    "shipping-guide.pdf": "Acme Retail – Shipping Guide",
    "product-care.docx": "Acme Product Care and Size Guide",
    "store-faq.md": "Acme Stores – Frequently Asked Questions",
    "warranty.txt": "Acme Warranty Terms",
}


def grouped(doc: str) -> OrderedDict[tuple[int | None, str], list[Fact]]:
    out: OrderedDict[tuple[int | None, str], list[Fact]] = OrderedDict()
    for f in FACTS:
        if f.doc == doc:
            out.setdefault((f.page, f.section), []).append(f)
    return out


def write_pdf(doc: str) -> None:
    pages: dict[int, list[tuple[str, int, str]]] = {}
    for (page, section), facts in grouped(doc).items():
        assert page is not None
        lines = pages.setdefault(page, [("H", 18, TITLES[doc])] if page == 1 else [])
        lines.append(("H", 13, section))
        for line in textwrap.wrap(" ".join(f.sentence for f in facts), 95):
            lines.append(("P", 10, line))
    (HERE / doc).write_bytes(build_pdf([pages[p] for p in sorted(pages)]))


def write_docx(doc: str) -> None:
    import docx

    d = docx.Document()
    d.add_heading(TITLES[doc], level=1)
    for (_, section), facts in grouped(doc).items():
        d.add_heading(section, level=2)
        d.add_paragraph(" ".join(f.sentence for f in facts))
    d.core_properties.author = "Acme Retail"
    d.save(str(HERE / doc))


def write_md(doc: str) -> None:
    parts = [f"# {TITLES[doc]}\n"]
    for (_, section), facts in grouped(doc).items():
        parts.append(f"## {section}\n\n" + " ".join(f.sentence for f in facts) + "\n")
    (HERE / doc).write_text("\n".join(parts), encoding="utf-8")


def write_txt(doc: str) -> None:
    parts = [TITLES[doc]]
    for facts in grouped(doc).values():
        parts.append(" ".join(f.sentence for f in facts))
    (HERE / doc).write_text("\n\n".join(parts) + "\n", encoding="utf-8")


def main() -> None:
    for doc in TITLES:
        {".pdf": write_pdf, ".docx": write_docx, ".md": write_md, ".txt": write_txt}[
            Path(doc).suffix
        ](doc)
    with (HERE / "questions.jsonl").open("w", encoding="utf-8") as out:
        for i, f in enumerate(FACTS, start=1):
            row = {
                "id": f"kb1-a{i:03d}",
                "type": "answerable",
                "question": f.question,
                "gold": {"file_name": f.doc, "page": f.page, "section": f.section},
                "evidence": f.evidence,
                "answer": f.sentence,
            }
            out.write(json.dumps(row, ensure_ascii=False) + "\n")
    print(f"wrote {len(TITLES)} documents and {len(FACTS)} questions to {HERE}")


if __name__ == "__main__":
    main()
